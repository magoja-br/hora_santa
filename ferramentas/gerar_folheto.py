#!/usr/bin/env python3
"""Gera o folheto digital da Hora Santa (site HTML + PDF A5 + Word) a partir de Markdown.

Uso:
    python gerar_folheto.py <pasta-do-mes>          # ex.: .../hora-santa/2026-10-02
    python gerar_folheto.py <pasta-do-mes> --sem-pdf
    python gerar_folheto.py --indice <pasta-raiz>   # só a página inicial e recursos

A pasta do mês deve conter `folheto.md` (para os fiéis) e, opcionalmente,
`organizacao.md` (para a equipe). Para cada um são gerados um .html
autocontido (CSS embutido, bom para enviar por WhatsApp) e um .pdf A5,
impresso pelo Microsoft Edge ou Google Chrome em modo headless.
O folheto ganha também um .docx A5, editável no Word (precisa do pacote
`python-docx`; sem ele, o .docx é pulado).
Depois o script refaz o `index.html` da pasta-mãe (arquivo de todos os meses).

Convenções do Markdown (além do Markdown comum):
    ---                         cabeçalho YAML simples (chave: valor) no topo
    ## 18h30 · Exposição        a hora no início do título vira um selo
    ! texto                     rubrica (instrução, em vermelho itálico)
    A. / L. / M. / P. texto     fala do Animador, Leitor, Ministro, Presidente
    T. / R. texto               resposta de Todos (em negrito)
    V. texto                    versículo (quem conduz)
    > [!Oração] Título          bloco de oração (cada linha "> ..." é um verso)
    > [!Nota] Título            caixa de observação
    ===pagina===                quebra de página no PDF
"""
import datetime, glob, html, os, re, shutil, subprocess, sys, tempfile

import markdown

AQUI = os.path.dirname(os.path.abspath(__file__))
CSS = next(c for c in (os.path.join(AQUI, "..", "assets", "folheto.css"),
                       os.path.join(AQUI, "folheto.css")) if os.path.exists(c))

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
QUEM = {"A": "Animador", "L": "Leitor", "M": "Ministro", "P": "Presidente",
        "T": "Todos", "R": "Todos", "V": ""}

EMBLEMA = """<svg class="emblema" viewBox="0 0 120 130" aria-hidden="true">
  <g fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round">
    <path d="M60 44 C60 44 57 30 50 24 M60 44 C60 44 63 30 70 24 M60 40 C60 40 60 26 60 18"/>
    <path d="M60 4 V24 M52 11 H68" stroke-width="3"/>
  </g>
  <path d="M60 118 C28 94 12 78 12 58 C12 44 23 34 36 34 C46 34 55 40 60 48 C65 40 74 34 84 34 C97 34 108 44 108 58 C108 78 92 94 60 118 Z"
        fill="currentColor"/>
  <g fill="none" stroke="#fbf7f0" stroke-width="2.4" stroke-linecap="round">
    <path d="M24 64 Q60 50 96 64" stroke-dasharray="6 5"/>
    <path d="M26 70 Q60 84 94 70" stroke-dasharray="6 5"/>
  </g>
  <path d="M63 92 l6 -8" stroke="#fbf7f0" stroke-width="3" stroke-linecap="round"/>
</svg>"""


def ler(caminho):
    t = open(caminho, encoding="utf-8").read()
    meta = {}
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", t, re.S)
    if m:
        for linha in m.group(1).splitlines():
            if ":" in linha:
                k, v = linha.split(":", 1)
                meta[k.strip()] = v.strip().strip('"').strip("'")
        t = t[m.end():]
    return meta, t


def data_extenso(iso):
    d = datetime.date.fromisoformat(iso)
    dia = "1º" if d.day == 1 else d.day
    return f"{dia} de {MESES[d.month - 1]} de {d.year}"


def converter(texto):
    """Traduz as convenções do folheto para Markdown + HTML (md_in_html)."""
    out, linhas, i = [], texto.splitlines(), 0
    while i < len(linhas):
        ln = linhas[i]
        s = ln.strip()
        m = re.match(r"^>\s*\[!(\w+)\]\s*(.*)$", s)
        if m:  # bloco de oração ou nota
            tipo, tit = m.group(1).lower(), m.group(2)
            classe = {"oração": "oracao", "oracao": "oracao"}.get(tipo, "nota")
            versos = []
            i += 1
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                versos.append(linhas[i].strip()[1:].strip())
                i += 1
            if classe == "oracao" and len(versos) <= 6:
                classe += " curta"  # oração curta não se divide entre páginas
            corpo = "  \n".join(v for v in versos)
            titulo = f'<p class="tit">{tit}</p>\n\n' if tit else ""
            out += [f'<div class="{classe}" markdown="1">', "", titulo + corpo, "", "</div>", ""]
            continue
        if s.startswith("- [ ]") or s.startswith("- [x]"):  # lista de compromisso
            itens = []
            while i < len(linhas) and re.match(r"^\s*- \[[ x]\]", linhas[i]):
                itens.append(re.sub(r"^\s*- \[( |x)\]", lambda m: "- " + ("☑" if m.group(1) == "x" else "☐"), linhas[i]))
                i += 1
            out += ['<div class="check" markdown="1">', ""] + itens + ["", "</div>", ""]
            continue
        if s == "===pagina===":
            out += ['<div class="quebra"></div>', ""]
        elif s.startswith("! "):
            out += ['<div class="rubrica" markdown="1">', "", s[2:], "", "</div>", ""]
        elif re.match(r"^([ALMPTRV])\.\s+", s):
            q, resto = s[0], s[2:].strip()
            classe = "fala todos" if q in "TR" else "fala"
            rot = f'<span class="quem" title="{QUEM[q]}">{q}.</span> '
            out += [f'<div class="{classe}" markdown="1">', "", rot + resto, "", "</div>", ""]
        else:
            m = re.match(r"^(#{2,3})\s+(\d{1,2}h\d{0,2})\s*[·—–-]\s*(.+)$", s)
            if m:
                out.append(f'{m.group(1)} <span class="hora">{m.group(2)}</span> {m.group(3)}')
            else:
                out.append(ln)
        i += 1
    return "\n".join(out)


def pagina(meta, corpo_html, css, doc, irmaos):
    tit = meta.get("titulo", "Hora Santa")
    data = data_extenso(meta["data"]) if meta.get("data") else ""
    prog = [p.split("|", 1) for p in meta.get("programacao", "").split(";") if "|" in p]
    prog_html = "".join(f"<li><b>{h.strip()}</b> {html.escape(o.strip())}</li>" for h, o in prog)
    links = "".join(f'<a href="{h}">{t}</a>' for h, t in irmaos)
    capa = f"""
<header class="capa">
  <p class="sobre">{html.escape(meta.get('sobretitulo', 'Primeira sexta-feira do mês'))}</p>
  {EMBLEMA}
  <h1>{html.escape(tit)}</h1>
  <p class="data">{data}</p>
  <p class="tema">{html.escape(meta.get('tema', ''))}</p>
  <p class="versiculo">{html.escape(meta.get('versiculo', ''))}</p>
  {f'<ul class="programa">{prog_html}</ul>' if prog_html else ''}
  <p class="paroquia">{html.escape(meta.get('paroquia', ''))}</p>
</header>"""
    return f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(tit)}{" — " + data if data else ""}</title>
<meta name="description" content="{html.escape(meta.get('tema', ''))}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:ital,wght@0,500;0,700;1,500&family=Source+Serif+4:ital,opsz,wght@0,8..60,400;0,8..60,600;1,8..60,400&display=swap" rel="stylesheet">
<style>{css}</style>
</head>
<body class="{doc}">
<nav class="barra">{links}</nav>
<main>
{capa}
<article>
{corpo_html}
</article>
<footer class="rodape">{html.escape(meta.get('rodape', 'Fonte: “No Coração de Jesus”, cap. 9 e Apêndice.'))}</footer>
</main>
</body>
</html>"""


def achar_navegador():
    candidatos = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]
    for c in candidatos:
        if os.path.exists(c):
            return c
    for n in ("msedge", "google-chrome", "chromium", "chrome"):
        if shutil.which(n):
            return shutil.which(n)
    return None


def imprimir_pdf(html_path, pdf_path):
    nav = achar_navegador()
    if not nav:
        print("  ! Edge/Chrome não encontrado: PDF não gerado.")
        return False
    perfil = tempfile.mkdtemp(prefix="folheto-")
    url = "file:///" + os.path.abspath(html_path).replace("\\", "/")
    cmd = [nav, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
           f"--user-data-dir={perfil}", "--virtual-time-budget=8000",
           f"--print-to-pdf={os.path.abspath(pdf_path)}", url]
    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    shutil.rmtree(perfil, ignore_errors=True)
    return os.path.exists(pdf_path)


def gerar_docx(meta, texto, caminho):
    """Gera o folheto em Word (A5), com as mesmas convenções do Markdown."""
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError:
        print("  ! python-docx não instalado: .docx não gerado (pip install python-docx).")
        return False
    RUBI, OURO, SUAVE = RGBColor(0x8F, 0x1D, 0x21), RGBColor(0xA9, 0x7B, 0x2C), RGBColor(0x6B, 0x5D, 0x52)
    FONTE = "Georgia"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(148), Mm(210)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(13), Mm(14), Mm(12), Mm(12)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONTE, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(4), 1.1
    for nome, tam in (("Heading 1", 15), ("Heading 2", 12)):
        h = doc.styles[nome]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = FONTE, Pt(tam), True, RUBI
        h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # usa a fonte, não o tema
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12 if tam > 12 else 8), Pt(4)
        h.paragraph_format.keep_with_next = True

    def inline(p, s, cor=None, italico=False, negrito=False):
        """Escreve **negrito** e *itálico* do Markdown como runs."""
        for parte in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", s):
            if not parte:
                continue
            b, i = negrito, italico
            if parte.startswith("**") and parte.endswith("**"):
                parte, b = parte[2:-2], True
            elif parte.startswith("*") and parte.endswith("*") and len(parte) > 1:
                parte, i = parte[1:-1], not italico
            r = p.add_run(parte)
            r.bold, r.italic = b, i
            if cor is not None:
                r.font.color.rgb = cor
        return p

    def paragrafo(s="", alinhar=None, tam=None, cor=None, italico=False, negrito=False, depois=None):
        p = inline(doc.add_paragraph(), s, cor, italico, negrito)
        if alinhar:
            p.alignment = alinhar
        if tam:
            for r in p.runs:
                r.font.size = Pt(tam)
        if depois is not None:
            p.paragraph_format.space_after = Pt(depois)
        return p

    def caixa(titulo, versos, oracao):
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cel = t.cell(0, 0)
        tcpr = cel._tc.get_or_add_tcPr()
        fundo = OxmlElement("w:shd")
        fundo.set(qn("w:val"), "clear"), fundo.set(qn("w:color"), "auto"), fundo.set(qn("w:fill"), "F4ECDF")
        tcpr.append(fundo)
        bordas = OxmlElement("w:tcBorders")
        for lado in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "12" if lado == "left" else "4")
            b.set(qn("w:color"), "8F1D21" if lado == "left" else "E4D8C6")
            bordas.append(b)
        tcpr.append(bordas)
        p = cel.paragraphs[0]
        if titulo:
            inline(p, titulo, RUBI, negrito=True)
            p = cel.add_paragraph()
        for n, v in enumerate(versos):
            if n:
                p = cel.add_paragraph()
            m = re.match(r"^([VR])\.\s+(.*)$", v)
            if m:
                r = p.add_run(m.group(1) + ". ")
                r.bold, r.font.color.rgb = True, RUBI
                v = m.group(2)
            inline(p, v)
            p.paragraph_format.space_after = Pt(2 if oracao else 3)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # capa
    paragrafo(meta.get("sobretitulo", "Primeira sexta-feira do mês").upper(), WD_ALIGN_PARAGRAPH.CENTER, 9, OURO, depois=18).paragraph_format.space_before = Pt(36)
    paragrafo("♥", WD_ALIGN_PARAGRAPH.CENTER, 34, RUBI, depois=6)
    paragrafo(meta.get("titulo", "Hora Santa"), WD_ALIGN_PARAGRAPH.CENTER, 30, RUBI, negrito=True, depois=4)
    if meta.get("data"):
        paragrafo(data_extenso(meta["data"]), WD_ALIGN_PARAGRAPH.CENTER, 12, depois=14)
    if meta.get("tema"):
        paragrafo(meta["tema"], WD_ALIGN_PARAGRAPH.CENTER, 15, RUBI, italico=True, depois=2)
    if meta.get("versiculo"):
        paragrafo(meta["versiculo"], WD_ALIGN_PARAGRAPH.CENTER, 9.5, SUAVE, depois=18)
    for item in meta.get("programacao", "").split(";"):
        if "|" in item:
            hora, o = item.split("|", 1)
            p = paragrafo("", WD_ALIGN_PARAGRAPH.CENTER, depois=2)
            r = p.add_run(hora.strip() + "  ")
            r.bold, r.font.color.rgb = True, RUBI
            p.add_run(o.strip())
    if meta.get("paroquia"):
        paragrafo(meta["paroquia"], WD_ALIGN_PARAGRAPH.CENTER, 9.5, SUAVE, depois=0).paragraph_format.space_before = Pt(24)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    linhas, i, lista = texto.splitlines(), 0, 0
    while i < len(linhas):
        s = linhas[i].strip()
        m = re.match(r"^>\s*\[!(\w+)\]\s*(.*)$", s)
        if m:
            versos = []
            i += 1
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                versos.append(linhas[i].strip()[1:].strip())
                i += 1
            caixa(m.group(2), versos, m.group(1).lower() in ("oração", "oracao"))
            continue
        i += 1
        if not s:
            lista = 0
            continue
        if s == "===pagina===":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("#"):
            nivel = 1 if s.startswith("## ") else 2
            doc.add_heading(re.sub(r"^#+\s*", "", s), level=nivel)
        elif s.startswith("! "):
            paragrafo(s[2:], cor=RUBI, italico=True, tam=9.5)
        elif re.match(r"^([ALMPTRV])\.\s+", s):
            p = doc.add_paragraph()
            r = p.add_run(s[:2] + " ")
            r.bold, r.font.color.rgb = True, RUBI
            inline(p, s[2:].strip(), negrito=s[0] in "TR")
        elif re.match(r"^- \[[ x]\]\s+", s):
            paragrafo(("☑ " if s[3] == "x" else "☐ ") + s[6:].strip())
        elif s.startswith("- "):
            p = paragrafo("• " + s[2:])
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Mm(5), Mm(-3.5)
        elif re.match(r"^\d+\.\s+", s):
            lista += 1
            p = paragrafo(f"{lista}. " + re.sub(r"^\d+\.\s+", "", s))
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Mm(5), Mm(-4)
        else:
            paragrafo(s)

    if meta.get("rodape"):
        paragrafo(meta["rodape"], WD_ALIGN_PARAGRAPH.CENTER, 8.5, SUAVE, italico=True).paragraph_format.space_before = Pt(18)
    # número da página no rodapé
    rp = sec.footer.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = rp.add_run()
        r.font.size, r.font.color.rgb = Pt(8.5), SUAVE
        if tipo:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
        else:
            f = OxmlElement("w:instrText")
            f.set(qn("xml:space"), "preserve")
            f.text = txt
        r._r.append(f)
    sec.different_first_page_header_footer = True
    doc.core_properties.title = f"{meta.get('titulo', 'Hora Santa')} — {data_extenso(meta['data'])}" if meta.get("data") else meta.get("titulo", "Hora Santa")
    doc.core_properties.author = meta.get("paroquia", "")
    doc.save(caminho)
    return True


def primeiras_sextas(n, desde=None):
    d = desde or datetime.date.today()
    ano, mes, out = d.year, d.month, []
    while len(out) < n:
        x = datetime.date(ano, mes, 1)
        while x.weekday() != 4:
            x += datetime.timedelta(1)
        if x >= d:
            out.append(x)
        ano, mes = (ano + 1, 1) if mes == 12 else (ano, mes + 1)
    return out


def paginas_raiz(raiz, css):
    """Converte as páginas .md da raiz (ex.: recursos.md) em .html."""
    for md in glob.glob(os.path.join(raiz, "*.md")):
        if os.path.basename(md).lower() == "readme.md":
            continue
        meta, texto = ler(md)
        corpo = markdown.markdown(converter(texto),
                                  extensions=["md_in_html", "tables", "sane_lists", "attr_list"])
        meta.setdefault("sobretitulo", "Hora Santa · primeira sexta-feira")
        h = md[:-3] + ".html"
        open(h, "w", encoding="utf-8").write(
            pagina(meta, corpo, css, "pagina", [("index.html", "Início")]))
        print("Página:", h)


def refazer_indice(raiz, css):
    paginas_raiz(raiz, css)
    feitos, itens = {}, []
    for pasta in sorted(glob.glob(os.path.join(raiz, "*", "folheto.md")), reverse=True):
        meta, _ = ler(pasta)
        nome = os.path.basename(os.path.dirname(pasta))
        feitos[nome] = meta
        data = data_extenso(meta["data"]) if meta.get("data") else nome
        itens.append(f"""<li><a href="{nome}/folheto.html"><b>{data}</b>
<span>{html.escape(meta.get('tema', ''))}</span></a>
<small><a href="{nome}/folheto.pdf">PDF</a>{f' · <a href="{nome}/folheto.docx">Word</a>' if os.path.exists(os.path.join(raiz, nome, "folheto.docx")) else ""} · <a href="{nome}/organizacao.html">organização</a></small></li>""")
    prox = []
    for d in primeiras_sextas(6):
        iso = d.isoformat()
        if iso in feitos:
            prox.append(f'<li><a href="{iso}/folheto.html"><b>{data_extenso(iso)}</b></a> '
                        f'— {html.escape(feitos[iso].get("tema", ""))}</li>')
        else:
            prox.append(f"<li><b>{data_extenso(iso)}</b> — <i>folheto em preparação</i></li>")
    links = [(f[:-3] + ".html", ler(os.path.join(raiz, f))[0].get("titulo", f[:-3]))
             for f in sorted(os.listdir(raiz)) if f.endswith(".md") and f.lower() != "readme.md"]
    barra = "".join(f'<a href="{h}">{html.escape(t)}</a>' for h, t in links)
    doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Horas Santas</title>
<meta name="description" content="Folhetos e organização da Hora Santa da primeira sexta-feira.">
<link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:ital,wght@0,500;0,700;1,500&family=Source+Serif+4:opsz,wght@8..60,400;8..60,600&display=swap" rel="stylesheet">
<style>{css}</style></head><body class="indice">
{f'<nav class="barra">{barra}</nav>' if barra else ''}<main>
<header class="capa curta"><p class="sobre">Primeira sexta-feira do mês</p>{EMBLEMA}<h1>Hora Santa</h1>
<p class="tema">"Ficai aqui e vigiai comigo" (Mt 26,38)</p></header>
<h2>Próximas datas</h2><ul class="proximas">{''.join(prox)}</ul>
<h2>Folhetos</h2><ul class="arquivo">{''.join(itens)}</ul></main></body></html>"""
    open(os.path.join(raiz, "index.html"), "w", encoding="utf-8").write(doc)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if sys.argv[1] == "--indice":  # só refaz index.html e páginas da raiz
        refazer_indice(os.path.abspath(sys.argv[2]), open(CSS, encoding="utf-8").read())
        return
    pasta = os.path.abspath(sys.argv[1])
    com_pdf = "--sem-pdf" not in sys.argv
    css = open(CSS, encoding="utf-8").read()
    docs = [d for d in ("folheto", "organizacao") if os.path.exists(os.path.join(pasta, d + ".md"))]
    rotulos = {"folheto": "Folheto", "organizacao": "Organização"}
    for doc in docs:
        meta, texto = ler(os.path.join(pasta, doc + ".md"))
        corpo = markdown.markdown(converter(texto),
                                  extensions=["md_in_html", "tables", "sane_lists", "attr_list"])
        irmaos = [(f"{d}.html", rotulos[d]) for d in docs]
        irmaos += [(f"{doc}.pdf", "Baixar PDF")]
        if doc == "folheto":
            w = os.path.join(pasta, "folheto.docx")
            print("Word:", w if gerar_docx(meta, texto, w) else "(não gerado)")
            if os.path.exists(w):
                irmaos += [("folheto.docx", "Baixar Word")]
        irmaos += [("../index.html", "Todos os meses")]
        h = os.path.join(pasta, doc + ".html")
        open(h, "w", encoding="utf-8").write(pagina(meta, corpo, css, doc, irmaos))
        print("HTML:", h)
        if com_pdf:
            p = os.path.join(pasta, doc + ".pdf")
            print("PDF: ", p if imprimir_pdf(h, p) else "(falhou)")
    refazer_indice(os.path.dirname(pasta), css)
    print("Índice:", os.path.join(os.path.dirname(pasta), "index.html"))


if __name__ == "__main__":
    main()
